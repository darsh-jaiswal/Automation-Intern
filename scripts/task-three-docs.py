"""Create Task 3 artifacts from observed answers; never manufacture provider results."""
from pathlib import Path
import csv
import json
import shutil
from html import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs' / 'task3'
DOCS.mkdir(parents=True, exist_ok=True)
markdown = (ROOT / 'docs' / 'Task_3_Ops_Judgement.md').read_text(encoding='utf-8')
answer = markdown.split('## 3a. A silent bug\n\n', 1)[1].split('\n\nThis is an analysis', 1)[0]
paragraphs = answer.replace('`', '').split('\n\n')

document = Document()
section = document.sections[0]
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.7)
normal = document.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
document.add_heading('Task 3a · A silent bug', 0)
document.add_paragraph('Darsh Jaiswal | Myntmore Systems & AI Automation Intern | Round 2')
for text in paragraphs:
    document.add_paragraph(text)
document.add_paragraph('Analysis of the supplied snippet. No production schema changes were executed.')
path = DOCS / 'Task_3a_Silent_Bug_Darsh.docx'
document.save(path)
(ROOT.parent / path.name).write_bytes(path.read_bytes())

styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='AssignmentBody', fontName='Helvetica', fontSize=10.5,
                          leading=15, spaceAfter=12, alignment=TA_LEFT))
story = [Paragraph('Task 3a: A silent bug', styles['Title']),
         Paragraph('Darsh Jaiswal | Myntmore Systems & AI Automation Intern | Round 2', styles['Normal']),
         Spacer(1, 8 * mm)]
story.extend(Paragraph(escape(text), styles['AssignmentBody']) for text in paragraphs)
story.append(Paragraph('Analysis of the supplied snippet. No production schema changes were executed.', styles['Normal']))
pdf = DOCS / 'Task_3a_Silent_Bug_Darsh.pdf'
SimpleDocTemplate(str(pdf), pagesize=A4, topMargin=17*mm, bottomMargin=17*mm,
                  leftMargin=18*mm, rightMargin=18*mm).build(story)
(ROOT.parent / pdf.name).write_bytes(pdf.read_bytes())

prompts = json.loads((ROOT / '.local/task3/prompts.json').read_text(encoding='utf-8'))
observations_path = DOCS / 'observations.json'
observations = json.loads(observations_path.read_text(encoding='utf-8')) if observations_path.exists() else {}
assessed = {(row['provider'], row['question_id']): row for row in observations.get('rows', [])}
screenshots = DOCS / 'screenshots'
screenshots.mkdir(exist_ok=True)
rows = []
for question in prompts['questions']:
    for provider in prompts['providers']:
        path = ROOT / f'.local/task3/raw/{provider.lower()}-q{question["id"]}.json'
        data = json.loads(path.read_text(encoding='utf-8')) if path.exists() else {}
        assessment = assessed.get((provider, question['id']), {})
        screenshot = data.get('screenshot', '')
        source_image = ROOT / '.local/task3/raw' / screenshot
        image_present = bool(screenshot) and source_image.is_file()
        complete = data.get('status') == 'complete' and image_present
        if image_present:
            shutil.copy2(source_image, screenshots / source_image.name)
        rows.append({'question_id': question['id'], 'ai': provider, 'question': question['prompt'],
                     'collection_status': 'complete' if complete else 'awaiting_access_or_capture',
                     'myntmore_appears': assessment.get('appears', 'Not observed') if complete else 'Not observed',
                     'position': assessment.get('position', '') if complete else 'Unavailable',
                     'who_appears_instead': assessment.get('others', '') if complete else 'No agency answer obtained',
                     'answer_summary': assessment.get('summary', '') if complete else assessment.get('summary', 'Not captured; browser access required.'),
                     'ai_cited_sources': '\n'.join(link['url'] for link in data.get('links', [])),
                     'source_note': assessment.get('source_note', ''),
                     'response_url': data.get('response_url', ''),
                     'screenshot': f'screenshots/{screenshot}' if image_present else '',
                     'observed_at': data.get('observed_at', ''),
                     'mode': data.get('mode', '')})
with (DOCS / 'AI_Search_Collection_Status.csv').open('w', encoding='utf-8-sig', newline='') as handle:
    writer = csv.DictWriter(handle, fieldnames=rows[0].keys())
    writer.writeheader()
    writer.writerows(rows)

complete_count = sum(row['collection_status'] == 'complete' for row in rows)
report = Document()
report.add_heading('Task 3b · AI Search Evidence · DRAFT', 0)
report.add_paragraph(f'{complete_count}/15 actual agency answers captured. Missing answers are unavailable, not negative Myntmore results.')
report.add_paragraph('Five identical prompts, fresh conversations, signed-out default interfaces. Position is order in an answer, not a stable ranking. Sources are recorded as the AI provided them, including inaccurate or search-wrapper URLs. Screenshots are bundled in the screenshots directory.')
for row in rows:
    report.add_heading(f'Q{row["question_id"]} — {row["ai"]}', 2)
    report.add_paragraph(row['question'])
    report.add_paragraph(f'Status: {row["collection_status"]}. Myntmore: {row["myntmore_appears"]}; position: {row["position"]}.')
    report.add_paragraph(f'Other names: {row["who_appears_instead"]}\nAnswer: {row["answer_summary"]}')
    report.add_paragraph(f'AI-cited sources:\n{row["ai_cited_sources"] or row["source_note"]}\nResponse URL: {row["response_url"] or "Not available"}\nScreenshot: {row["screenshot"] or "No answer screenshot"}')
    if row['source_note'] and row['ai_cited_sources']:
        report.add_paragraph(row['source_note'])
report.add_heading('Three website changes — provisional', 1)
for change in observations.get('website_changes', []):
    report.add_heading(f'{change["id"]}. {change["questions"]}', 2)
    report.add_paragraph(change['change'])
    report.add_paragraph(change['why'])
report.save(DOCS / 'Task_3_AI_Search_DRAFT.docx')

table_style = ParagraphStyle(name='EvidenceCell', fontName='Helvetica', fontSize=7, leading=9, wordWrap='CJK')
cell = lambda value: Paragraph(escape(str(value)).replace('\n', '<br/>'), table_style)
table_rows = [[cell(value) for value in ['#', 'AI', 'Question', 'Myntmore?', 'Who appears instead', 'Answer summary', 'AI-cited sources / notes', 'Screenshot']]]
for row in rows:
    table_rows.append([cell(value) for value in [row['question_id'], row['ai'], row['question'],
                      f'{row["myntmore_appears"]}; {row["position"]}', row['who_appears_instead'],
                      row['answer_summary'], row['ai_cited_sources'] or row['source_note'],
                      row['screenshot'] or 'No answer capture']])
table = Table(table_rows, colWidths=[16, 42, 117, 50, 95, 108, 174, 65], repeatRows=1)
table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e6ece5')),
                           ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#bac4b6')),
                           ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                           ('LEFTPADDING', (0, 0), (-1, -1), 5),
                           ('RIGHTPADDING', (0, 0), (-1, -1), 5)]))
evidence_pdf = DOCS / 'Task_3_AI_Search_DRAFT.pdf'
SimpleDocTemplate(str(evidence_pdf), pagesize=landscape(A4), topMargin=13*mm, bottomMargin=13*mm,
                  leftMargin=12*mm, rightMargin=12*mm).build([
                      Paragraph(f'Task 3b: AI search — DRAFT ({complete_count}/15 answers)', styles['Title']),
                      Paragraph('Unavailable means no agency answer was obtained, not that Myntmore was absent. See bundled full-resolution screenshots. Gemini: Flash-Lite; Perplexity: Search; signed out. Sources are AI-provided and not endorsements.', styles['Normal']),
                      Spacer(1, 5*mm), table])

one_page_style = ParagraphStyle(name='WrittenBody', fontName='Helvetica', fontSize=9.6, leading=12.5, spaceAfter=8)
written = [Paragraph('Task 3: Written answers — DRAFT', styles['Title']),
           Paragraph(f'Darsh Jaiswal | {complete_count}/15 AI answers captured; remaining evidence pending browser access.', styles['Normal']),
           Spacer(1, 4*mm), Paragraph('3a. A silent bug', styles['Heading2'])]
written.extend(Paragraph(escape(text), one_page_style) for text in paragraphs)
written.append(Paragraph('3b. First three website changes (provisional)', styles['Heading2']))
for change in observations.get('website_changes', []):
    written.append(Paragraph(escape(f'{change["id"]}. {change["change"]} Why: {change["why"]} Helps: {change["questions"]}.'), one_page_style))
written.append(Paragraph('AI-search table and original screenshots are separate evidence. Finalise the recommendations after all 15 runs. No production website or database was modified.', styles['Normal']))
SimpleDocTemplate(str(DOCS / 'Task_3_Written_Answers_DRAFT.pdf'), pagesize=A4,
                  topMargin=13*mm, bottomMargin=13*mm, leftMargin=16*mm, rightMargin=16*mm).build(written)
print(json.dumps({'bug_answer_words': len(answer.split()), 'observations_needed': len(rows),
                  'complete_captures': complete_count,
                  'status': 'Collection in progress; visibility cells are not fabricated.'}))
